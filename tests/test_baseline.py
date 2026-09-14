# Arnav Sahu
# 24BCE2976

"""Focused endpoint and recognizer tests using small generated real files."""

from io import BytesIO
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from PIL.TiffImagePlugin import IFDRational
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, NameObject
from pytest import approx

from app.main import app
from app.extractors import image
from app.scan import mask_financial_text


client = TestClient(app)


def _docx_bytes(text: str = "Visible report") -> bytes:
    """Create a small real DOCX fixture."""

    document = Document()
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _authored_docx_bytes() -> bytes:
    """Create a DOCX whose author is absent from visible text."""

    document = Document()
    document.add_paragraph("Written by the visible team")
    document.core_properties.author = "Hidden Author"
    document.core_properties.last_modified_by = "Hidden Editor"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _hidden_text_docx_bytes() -> bytes:
    """Create a DOCX containing a literal Word hidden-text run."""

    source = _docx_bytes("Visible report")
    output = BytesIO()
    with ZipFile(BytesIO(source)) as source_archive, ZipFile(output, "w") as output_archive:
        for item in source_archive.infolist():
            data = source_archive.read(item.filename)
            if item.filename == "word/document.xml":
                xml = data.decode("utf-8")
                hidden_run = '<w:r><w:rPr><w:vanish/></w:rPr><w:t>SECRET</w:t></w:r>'
                data = xml.replace("</w:p>", hidden_run + "</w:p>").encode("utf-8")
            output_archive.writestr(item, data)
    return output.getvalue()


def _gps_jpeg_bytes() -> bytes:
    """Create a real JPEG containing device, timestamp, and GPS EXIF."""

    image = Image.new("RGB", (2, 2), "white")
    exif = Image.Exif()
    exif[271] = "Canon"
    exif[272] = "Test Camera"
    exif[36867] = "2024:01:02 03:04:05"
    exif[34853] = {1: "N", 2: (IFDRational(12, 1), IFDRational(34, 1), IFDRational(56, 1)), 3: "E", 4: (IFDRational(77, 1), IFDRational(8, 1), IFDRational(9, 1))}
    stream = BytesIO()
    image.save(stream, "JPEG", exif=exif)
    return stream.getvalue()


def _redacted_pdf_bytes() -> bytes:
    """Create a real PDF with text underneath a filled black rectangle."""

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 200 Td (SECRET) Tj ET 50 190 100 20 re 0 0 0 rg f")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _xlsx_bytes() -> bytes:
    """Create a workbook with a populated very-hidden sheet."""

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "PAN ABCDE1234F UPI user@ybl"
    hidden = workbook.create_sheet("PriorData")
    hidden.sheet_state = "veryHidden"
    hidden["A1"] = "historical customer data"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_very_hidden_sheet_is_reported() -> None:
    """A populated very-hidden worksheet creates a high-severity finding."""

    response = client.post("/scan/baseline", files={"file": ("evidence.xlsx", _xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200
    body = response.json()
    assert any(item["type"] == "veryHidden" for item in body["metadata"]["hidden_content"])
    assert "high_severity_very_hidden_data" in body["severity_flags"]


def test_image_gps_and_device_metadata_are_extracted() -> None:
    """Image EXIF GPS is converted to decimal degrees."""

    response = client.post("/scan/baseline", files={"file": ("photo.jpg", _gps_jpeg_bytes(), "image/jpeg")})
    assert response.status_code == 200
    body = response.json()
    assert body["metadata"]["gps"] == approx({"lat": 12.582222222222223, "lon": 77.13583333333333})
    assert body["metadata"]["device"] == "Canon Test Camera"


def test_heic_exif_support_is_reported_explicitly() -> None:
    """HEIC containers do not silently appear clean without a HEIF decoder."""

    result = image.extract(b"\x00\x00\x00\x18ftypheic\x00\x00\x00\x00")
    assert "heic_exif_extraction_unsupported" in result.severity_flags
    assert result.metadata.hidden_content[0]["type"] == "heic_exif_extraction_unsupported"


def test_docx_author_mismatch_is_flagged() -> None:
    """A core author absent from visible text is reported as a leak signal."""

    response = client.post("/scan/baseline", files={"file": ("evidence.docx", _authored_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200
    assert "author_visible_text_mismatch" in response.json()["severity_flags"]


def test_docx_hidden_text_run_is_reported() -> None:
    """Literal w:vanish text is reported when absent from visible text."""

    response = client.post("/scan/baseline", files={"file": ("evidence.docx", _hidden_text_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200
    body = response.json()
    assert "hidden_text_run" in body["severity_flags"]
    assert any(item["type"] == "hidden_text_run" and item["summary"] == "SECRET" for item in body["metadata"]["hidden_content"])


def test_pdf_redaction_failure_recovers_text() -> None:
    """Text below a filled dark rectangle is returned as a redaction failure."""

    response = client.post("/scan/baseline", files={"file": ("evidence.pdf", _redacted_pdf_bytes(), "application/pdf")})
    assert response.status_code == 200
    failures = response.json()["redaction_failures"]
    assert failures and failures[0]["page"] == 1 and "SECRET" in failures[0]["recovered_text"]


def test_financial_values_are_masked_for_display() -> None:
    """Financial values are masked before appearing in the public response."""

    financial_text = "PAN ABCDE1234F Aadhaar 2345 6789 0124 and UPI user@ybl"
    response = client.post("/scan/baseline", files={"file": ("evidence.docx", _docx_bytes(financial_text), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200
    body = response.json()
    assert any(item["entity_type"] == "FIN_PAN_CARD" for item in body["financial_findings"])
    assert "ABCDE1234F" not in response.text
    assert any(item["entity_type"] == "FIN_UPI_ID" for item in body["financial_findings"])
    assert any(item["entity_type"] == "FIN_AADHAAR" for item in body["financial_findings"])
    assert "2345 6789 0124" not in response.text
    assert all(value not in response.text for value in ("ABCDE1234F", "2345 6789 0124", "user@ybl"))


def test_random_number_is_not_aadhaar() -> None:
    """A generic twelve-digit number is not reported as Aadhaar."""

    response = client.post("/scan/baseline", files={"file": ("numbers.docx", _docx_bytes("Reference 123456789012"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200
    assert not any(item["entity_type"] == "FIN_AADHAAR" for item in response.json()["financial_findings"])


def test_unsupported_and_corrupt_files() -> None:
    """Unsupported files get 415 and corrupt supported files get a graceful 200."""

    unsupported = client.post("/scan/baseline", files={"file": ("evidence.txt", b"text", "text/plain")})
    corrupt = client.post("/scan/baseline", files={"file": ("evidence.pdf", b"not a pdf", "application/pdf")})
    assert unsupported.status_code == 415
    assert corrupt.status_code == 200
    assert corrupt.json()["error"]


def test_upload_size_limit() -> None:
    """Oversized uploads are rejected before scanning."""

    response = client.post("/scan/baseline", files={"file": ("evidence.pdf", b"0" * (25 * 1024 * 1024 + 1), "application/pdf")})
    assert response.status_code == 413
    assert "25 MB" in response.json()["detail"]


def test_mask_shape() -> None:
    """Display masking retains only the final four digits."""

    assert mask_financial_text("1234 5678 9012") == "XXXX XXXX 9012"