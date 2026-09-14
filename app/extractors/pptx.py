# Arnav Sahu
# 24BCE2976

"""Presentation properties, slide text, speaker notes, and hidden slides."""

from io import BytesIO

from pptx import Presentation

from app.models import ExtractionResult


def extract(content: bytes) -> ExtractionResult:
    """Extract presentation properties, visible text, notes, and hidden slides."""

    presentation = Presentation(BytesIO(content))
    result = ExtractionResult()
    properties = presentation.core_properties
    result.metadata.author = properties.author or None
    result.metadata.last_modified_by = properties.last_modified_by or None
    result.metadata.timestamps = {"created": properties.created.isoformat() if properties.created else None, "modified": properties.modified.isoformat() if properties.modified else None}
    text: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        text.extend(shape.text for shape in slide.shapes if hasattr(shape, "text"))
        notes = getattr(slide, "notes_slide", None)
        if notes and notes.notes_text_frame and notes.notes_text_frame.text.strip():
            notes_text = notes.notes_text_frame.text.strip()
            text.append(notes_text)
            result.metadata.hidden_content.append({"type": "speaker_notes", "location": f"slide:{number}", "summary": notes_text})
        if slide._element.get("show") in {"0", "false"}:
            result.metadata.hidden_content.append({"type": "hidden_slide", "location": f"slide:{number}", "summary": "Slide is marked hidden in slide show settings."})
    result.text = "\n".join(text)
    return result