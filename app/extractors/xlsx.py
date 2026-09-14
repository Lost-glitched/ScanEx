# Arnav Sahu
# 24BCE2976

"""Spreadsheet metadata, hidden-sheet, defined-name, and comment extraction."""

from io import BytesIO

from openpyxl import load_workbook

from app.models import ExtractionResult


def extract(content: bytes) -> ExtractionResult:
    """Extract workbook properties, every sheet state, names, comments, and cell text."""

    workbook = load_workbook(BytesIO(content), read_only=False, data_only=False)
    result = ExtractionResult()
    properties = workbook.properties
    result.metadata.author = properties.creator or None
    result.metadata.last_modified_by = properties.lastModifiedBy or None
    result.metadata.timestamps = {"created": properties.created.isoformat() if properties.created else None, "modified": properties.modified.isoformat() if properties.modified else None}
    text: list[str] = []
    for sheet in workbook.worksheets:
        has_data = sheet.max_row > 0 and sheet.max_column > 0 and any(cell.value is not None for row in sheet.iter_rows() for cell in row)
        text.extend(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None)
        if sheet.sheet_state in {"hidden", "veryHidden"}:
            result.metadata.hidden_content.append({"type": sheet.sheet_state, "location": f"sheet:{sheet.title}", "summary": f"{sheet.title} ({'contains data' if has_data else 'empty'})"})
            if sheet.sheet_state == "veryHidden" and has_data:
                result.severity_flags.append("high_severity_very_hidden_data")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment:
                    result.metadata.hidden_content.append({"type": "cell_comment", "location": f"{sheet.title}!{cell.coordinate}", "summary": cell.comment.text})
    result.text = "\n".join(text + [str(name) for name in workbook.defined_names])
    for name in workbook.defined_names:
        result.metadata.hidden_content.append({"type": "defined_name", "location": "workbook.defined_names", "summary": str(name)})
    workbook.close()
    return result